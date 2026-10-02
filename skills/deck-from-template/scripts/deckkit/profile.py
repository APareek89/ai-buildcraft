from __future__ import annotations

import re
import tempfile
import zipfile
from collections import defaultdict
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from .capacity import measured_capacity, resolve_font
from .models import (
    SlideSize,
    SlideTypeProfile,
    SlotProfile,
    TemplateProfile,
    ThemeProfile,
)

SlideTypeId = str

EMU = 914400
REQUIRED_TYPES: list[SlideTypeId] = [
    "title",
    "section",
    "big_numbers",
    "cards_2",
    "cards_3",
    "cards_4",
    "chart_insight",
    "table",
    "ranked_list",
    "flow",
    "closing",
]

LABELS = {
    "title": "Opening statement",
    "section": "Section divider",
    "big_numbers": "Three or four anchor numbers",
    "cards_2": "Two parallel points",
    "cards_3": "Three parallel points",
    "cards_4": "Four parallel points",
    "chart_insight": "Chart with interpretation",
    "table": "Evidence table",
    "ranked_list": "Ranked priorities",
    "flow": "Stages or relationships",
    "closing": "Closing decision",
}


def _presentation_path(path: Path) -> tuple[Path, tempfile.TemporaryDirectory[str] | None]:
    if path.suffix.lower() != ".potx":
        return path, None
    temp = tempfile.TemporaryDirectory(prefix="deckkit-potx-")
    output = Path(temp.name) / "template.pptx"
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as dest:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(
                    b"application/vnd.openxmlformats-officedocument.presentationml.template.main+xml",
                    b"application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml",
                )
            dest.writestr(item, data)
    return output, temp


def open_presentation(path: Path) -> Presentation:
    normalised, temp = _presentation_path(path)
    presentation = Presentation(str(normalised))
    presentation._deckkit_temp = temp  # keep the extracted copy alive
    return presentation


def _theme_from_package(path: Path) -> tuple[str, str, dict[str, str]]:
    defaults = {
        "dk1": "171717",
        "lt1": "FFFFFF",
        "dk2": "334155",
        "lt2": "F8FAFC",
        "accent1": "2563EB",
        "accent2": "14B8A6",
        "accent3": "F59E0B",
        "accent4": "8B5CF6",
        "accent5": "EC4899",
        "accent6": "64748B",
        "hlink": "2563EB",
        "folHlink": "7C3AED",
    }
    try:
        with zipfile.ZipFile(path) as archive:
            theme_name = next(name for name in archive.namelist() if name.startswith("ppt/theme/theme"))
            xml = archive.read(theme_name).decode("utf-8", errors="replace")
        major = re.search(r"<a:majorFont>.*?<a:latin[^>]+typeface=\"([^\"]*)\"", xml, re.DOTALL)
        minor = re.search(r"<a:minorFont>.*?<a:latin[^>]+typeface=\"([^\"]*)\"", xml, re.DOTALL)
        for key in defaults:
            match = re.search(
                rf"<a:{re.escape(key)}>.*?<(?:a:srgbClr|a:sysClr)[^>]+(?:val|lastClr)=\"([0-9A-Fa-f]{{6}})\"",
                xml,
                re.DOTALL,
            )
            if match:
                defaults[key] = match.group(1).upper()
        return (major.group(1) if major and major.group(1) else "Arial"), (
            minor.group(1) if minor and minor.group(1) else "Arial"
        ), defaults
    except (OSError, StopIteration, zipfile.BadZipFile):
        return "Arial", "Arial", defaults


def _font_size_from_element(element: Any, default: float) -> float:
    for node in element.iter():
        if node.tag.endswith("}defRPr") or node.tag.endswith("}rPr"):
            size = node.get("sz")
            if size:
                return max(8.0, int(size) / 100)
    return default


def _font_style_from_element(element: Any) -> tuple[str, str | None]:
    weight = "regular"
    color = None
    for node in element.iter():
        if node.tag.endswith("}defRPr") or node.tag.endswith("}rPr"):
            if node.get("b") in {"1", "true"}:
                weight = "bold"
            for child in node.iter():
                if child.tag.endswith("}srgbClr") and child.get("val"):
                    color = str(child.get("val")).upper()
                    break
    return weight, color


def _placeholder_default_font(ph: Any) -> float:
    kind = ph.placeholder_format.type
    if kind in {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}:
        fallback = 32.0
    elif kind == PP_PLACEHOLDER.SUBTITLE:
        fallback = 20.0
    else:
        fallback = 18.0
    direct = _font_size_from_element(ph.element, 0)
    if direct:
        return direct
    try:
        master = ph.part.slide_layout.slide_master.element
        style_name = "titleStyle" if kind in {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE} else "bodyStyle"
        for node in master.iter():
            if node.tag.endswith("}" + style_name):
                return _font_size_from_element(node, fallback)
    except (AttributeError, TypeError):
        return fallback
    return fallback


def _font_steps(font_pt: float) -> list[float]:
    whitelist = [40, 36, 32, 28, 24, 22, 20, 18, 16, 14, 12, 10]
    steps = [float(size) for size in whitelist if size <= font_pt]
    if font_pt not in steps:
        steps.insert(0, float(font_pt))
    return steps[:3]


def _slot(
    name: str,
    box: tuple[int, int, int, int],
    font_pt: float,
    font: str,
    ph_idx: int | None = None,
    font_weight: str = "regular",
    font_color: str | None = None,
) -> SlotProfile:
    left, top, width, height = box
    return SlotProfile(
        name=name,
        ph_idx=ph_idx,
        max_chars=measured_capacity(
            typeface=font, font_pt=font_pt, width_emu=width, height_emu=height
        ),
        font_pt=font_pt,
        font_steps=_font_steps(font_pt),
        font_weight=font_weight,
        font_color=font_color,
        left=left,
        top=top,
        width=width,
        height=height,
    )


def _component_slots(type_id: SlideTypeId, w: int, h: int, font: str) -> list[SlotProfile]:
    margin = int(w * 0.065)
    content_w = w - 2 * margin
    headline = _slot("headline", (margin, int(h * 0.075), content_w, int(h * 0.15)), 30, font)
    if type_id == "title":
        return [
            _slot("kicker", (margin, int(h * 0.21), content_w, int(h * 0.07)), 14, font),
            _slot("headline", (margin, int(h * 0.31), int(w * 0.72), int(h * 0.29)), 38, font),
            _slot("subhead", (margin, int(h * 0.64), int(w * 0.68), int(h * 0.13)), 18, font),
        ]
    if type_id == "section":
        return [
            _slot("eyebrow", (margin, int(h * 0.28), content_w, int(h * 0.07)), 14, font),
            _slot("headline", (margin, int(h * 0.39), int(w * 0.78), int(h * 0.20)), 36, font),
            _slot("body", (margin, int(h * 0.64), int(w * 0.65), int(h * 0.12)), 18, font),
        ]
    if type_id == "closing":
        return [
            _slot("headline", (margin, int(h * 0.24), int(w * 0.78), int(h * 0.24)), 38, font),
            _slot("body", (margin, int(h * 0.53), int(w * 0.68), int(h * 0.14)), 20, font),
            _slot("cta", (margin, int(h * 0.73), int(w * 0.55), int(h * 0.08)), 15, font),
        ]
    if type_id == "big_numbers":
        slots = [headline]
        gap = int(w * 0.018)
        col_w = int((content_w - gap * 3) / 4)
        for index in range(4):
            x = margin + index * (col_w + gap)
            slots.extend(
                [
                    _slot(f"number_{index + 1}", (x, int(h * 0.37), col_w, int(h * 0.16)), 34, font),
                    _slot(f"label_{index + 1}", (x, int(h * 0.57), col_w, int(h * 0.15)), 16, font),
                ]
            )
        return slots
    if type_id.startswith("cards_"):
        count = int(type_id[-1])
        slots = [headline]
        gap = int(w * 0.018)
        col_w = int((content_w - gap * (count - 1)) / count)
        for index in range(count):
            x = margin + index * (col_w + gap)
            slots.extend(
                [
                    _slot(f"card_{index + 1}_title", (x, int(h * 0.30), col_w, int(h * 0.12)), 18, font),
                    _slot(f"card_{index + 1}_body", (x, int(h * 0.44), col_w, int(h * 0.34)), 14, font),
                ]
            )
        return slots
    if type_id == "chart_insight":
        return [
            headline,
            _slot("subhead", (margin, int(h * 0.245), content_w, int(h * 0.06)), 16, font),
            _slot("chart", (margin, int(h * 0.34), int(content_w * 0.67), int(h * 0.48)), 12, font),
            _slot("insight_1_title", (int(w * 0.74), int(h * 0.39), int(w * 0.20), int(h * 0.10)), 17, font),
            _slot("insight_1_body", (int(w * 0.74), int(h * 0.51), int(w * 0.20), int(h * 0.20)), 14, font),
            _slot("source", (margin, int(h * 0.86), content_w, int(h * 0.04)), 10, font),
        ]
    if type_id == "table":
        return [
            headline,
            _slot("table", (margin, int(h * 0.29), content_w, int(h * 0.51)), 13, font),
            _slot("source", (margin, int(h * 0.84), content_w, int(h * 0.04)), 10, font),
        ]
    if type_id == "ranked_list":
        slots = [headline]
        for index in range(1, 7):
            y = int(h * (0.29 + (index - 1) * 0.095))
            slots.extend(
                [
                    _slot(f"item_{index}_title", (margin, y, int(content_w * 0.35), int(h * 0.07)), 16, font),
                    _slot(f"item_{index}_body", (int(w * 0.42), y, int(content_w * 0.60), int(h * 0.07)), 14, font),
                ]
            )
        return slots
    if type_id == "flow":
        return [
            headline,
            _slot("flow", (margin, int(h * 0.31), content_w, int(h * 0.45)), 14, font),
            _slot("source", (margin, int(h * 0.84), content_w, int(h * 0.04)), 10, font),
        ]
    raise ValueError(type_id)


def _guess_type(name: str, placeholders: int) -> SlideTypeId | None:
    lower = name.lower()
    patterns: list[tuple[str, SlideTypeId]] = [
        (r"title|cover", "title"),
        (r"section|divider|chapter", "section"),
        (r"closing|thank|end", "closing"),
        (r"chart|graph|data", "chart_insight"),
        (r"table", "table"),
        (r"flow|process|timeline|journey", "flow"),
        (r"rank|agenda|list", "ranked_list"),
        (r"four|4", "cards_4"),
        (r"three|3", "cards_3"),
        (r"two|2|comparison", "cards_2"),
        (r"number|metric|kpi", "big_numbers"),
    ]
    for pattern, type_id in patterns:
        if re.search(pattern, lower):
            return type_id
    if placeholders >= 9:
        return "cards_4"
    if placeholders >= 7:
        return "cards_3"
    if placeholders >= 5:
        return "cards_2"
    return None


def _geometry_ok(slots: list[SlotProfile], w: int, h: int) -> bool:
    """Reject a layout whose repurposed boxes would fail the deck lint anyway.

    Template layouts are only borrowed when their placeholder geometry already
    respects the safe margin and keeps the slots apart; otherwise the measured
    component grid is the safer home for this slide type.
    """
    margin = int(0.4 * EMU)
    pad = int(0.075 * EMU)
    boxes = []
    for slot in slots:
        left, top = slot.left, slot.top
        right, bottom = left + slot.width, top + slot.height
        if left < margin or top < margin or w - right < margin or h - bottom < margin:
            return False
        boxes.append((left - pad, top - pad, right + pad, bottom + pad))
    for index, first in enumerate(boxes):
        for second in boxes[index + 1 :]:
            if min(first[2], second[2]) > max(first[0], second[0]) and min(
                first[3], second[3]
            ) > max(first[1], second[1]):
                return False
    return True


def _covers_type(slots: list[SlotProfile], type_id: SlideTypeId, w: int, h: int, font: str) -> bool:
    """A layout may only claim a slide type when it can hold every text slot that type needs.

    Without this, a two-placeholder "Title and Content" layout happily claims
    cards_2 and the second card is dropped on the floor at render time.
    """
    expected = {
        slot.name
        for slot in _component_slots(type_id, w, h, font)
        if slot.name not in {"chart", "table", "flow"}
    }
    return expected.issubset({slot.name for slot in slots})


def _layout_slots(layout: Any, type_id: SlideTypeId, font: str) -> list[SlotProfile]:
    placeholders = []
    for ph in layout.placeholders:
        kind = ph.placeholder_format.type
        if kind in {
            PP_PLACEHOLDER.DATE,
            PP_PLACEHOLDER.FOOTER,
            PP_PLACEHOLDER.SLIDE_NUMBER,
        }:
            continue
        if int(ph.width) <= 0 or int(ph.height) <= 0:
            continue
        placeholders.append(ph)
    placeholders.sort(key=lambda ph: (int(ph.top), int(ph.left)))
    expected = [slot.name for slot in _component_slots(type_id, layout.part.slide_width if hasattr(layout.part, "slide_width") else 12192000, layout.part.slide_height if hasattr(layout.part, "slide_height") else 6858000, font)]
    non_visual = [name for name in expected if name not in {"chart", "table", "flow"}]
    used: set[str] = set()
    result: list[SlotProfile] = []
    for ph in placeholders:
        kind = ph.placeholder_format.type
        if kind in {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE} and "headline" in non_visual:
            name = "headline"
        elif kind == PP_PLACEHOLDER.SUBTITLE and "subhead" in non_visual:
            name = "subhead"
        else:
            name = next((candidate for candidate in non_visual if candidate not in used), f"body_{len(used) + 1}")
        if name in used:
            name = f"{name}_{len(used) + 1}"
        used.add(name)
        font_pt = _placeholder_default_font(ph)
        weight, color = _font_style_from_element(ph.element)
        if weight == "regular" and color is None:
            try:
                master = ph.part.slide_layout.slide_master.element
                inherited_weight, inherited_color = _font_style_from_element(master)
                weight = inherited_weight
                color = inherited_color
            except (AttributeError, TypeError):
                pass
        result.append(
            _slot(
                name,
                (int(ph.left), int(ph.top), int(ph.width), int(ph.height)),
                font_pt,
                font,
                int(ph.placeholder_format.idx),
                weight,
                color,
            )
        )
    component = {slot.name: slot for slot in _component_slots(type_id, 12192000, 6858000, font)}
    for visual in ("chart", "table", "flow"):
        if visual in expected and visual not in used and visual in component:
            result.append(component[visual])
    return result


def _blank_layout_index(prs: Presentation) -> int:
    for index, layout in enumerate(prs.slide_layouts):
        if "blank" in layout.name.lower():
            return index
    counts = []
    for index, layout in enumerate(prs.slide_layouts):
        visible = sum(
            1
            for ph in layout.placeholders
            if ph.placeholder_format.type
            not in {PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER}
        )
        counts.append((visible, index))
    return min(counts)[1] if counts else 0


def _slide_signature(prs: Presentation, slide: Any) -> tuple[Any, ...]:
    items = []
    for shape in slide.shapes:
        items.append(
            (
                round(int(shape.left) / max(1, prs.slide_width), 1),
                round(int(shape.top) / max(1, prs.slide_height), 1),
                round(int(shape.width) / max(1, prs.slide_width), 1),
                round(int(shape.height) / max(1, prs.slide_height), 1),
                bool(getattr(shape, "has_text_frame", False)),
            )
        )
    return tuple(sorted(items))


def _shape_font_style(shape: Any) -> tuple[float, str, str | None]:
    if not getattr(shape, "has_text_frame", False):
        return _font_size_from_element(shape.element, 16), "regular", None
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            size = run.font.size.pt if run.font.size else _font_size_from_element(shape.element, 16)
            weight = "bold" if run.font.bold else "regular"
            try:
                color = str(run.font.color.rgb) if run.font.color.rgb else None
            except (AttributeError, TypeError):
                color = None
            return float(size), weight, color
    return _font_size_from_element(shape.element, 16), "regular", None


def _cluster_slide_profiles(
    prs: Presentation, font: str, blank_index: int
) -> dict[SlideTypeId, SlideTypeProfile]:
    clusters: dict[tuple[Any, ...], list[Any]] = defaultdict(list)
    for slide in prs.slides:
        clusters[_slide_signature(prs, slide)].append(slide)
    mapped: dict[SlideTypeId, SlideTypeProfile] = {}
    for slides in sorted(clusters.values(), key=len, reverse=True):
        representative = slides[0]
        text_shapes = [
            shape
            for shape in representative.shapes
            if getattr(shape, "has_text_frame", False)
            and int(shape.width) > 0
            and int(shape.height) > 0
        ]
        text_shapes.sort(key=lambda shape: (int(shape.top), int(shape.left)))
        combined = " ".join(shape.text for shape in text_shapes if shape.text)
        type_id = _guess_type(combined, len(text_shapes))
        if type_id is None:
            if len(text_shapes) <= 2 and "title" not in mapped:
                type_id = "title"
            elif len(text_shapes) >= 9:
                type_id = "cards_4"
            elif len(text_shapes) >= 7:
                type_id = "cards_3"
            elif len(text_shapes) >= 5:
                type_id = "cards_2"
            else:
                type_id = "ranked_list"
        if type_id in mapped or not text_shapes:
            continue
        expected = [
            slot.name
            for slot in _component_slots(type_id, int(prs.slide_width), int(prs.slide_height), font)
            if slot.name not in {"chart", "table", "flow"}
        ]
        slots = []
        for index, shape in enumerate(text_shapes):
            name = expected[index] if index < len(expected) else f"body_{index + 1}"
            font_pt, weight, color = _shape_font_style(shape)
            slots.append(
                _slot(
                    name,
                    (int(shape.left), int(shape.top), int(shape.width), int(shape.height)),
                    font_pt,
                    font,
                    None,
                    weight,
                    color,
                )
            )
        if type_id in {"chart_insight", "table", "flow"}:
            component = {
                slot.name: slot
                for slot in _component_slots(
                    type_id, int(prs.slide_width), int(prs.slide_height), font
                )
            }
            visual_name = {"chart_insight": "chart", "table": "table", "flow": "flow"}[
                type_id
            ]
            slots.append(component[visual_name])
        mapped[type_id] = SlideTypeProfile(
            id=type_id,
            label=LABELS[type_id],
            layout_index=blank_index,
            slots=slots,
            supports_chart=type_id == "chart_insight",
            origin="component",
        )
    return mapped


def extract_template_profile(path: Path) -> TemplateProfile:
    path = path.resolve()
    prs = open_presentation(path)
    major, minor, colors = _theme_from_package(path)
    font_available: dict[str, bool] = {}
    for name in {major, minor}:
        try:
            _, exact = resolve_font(name)
            font_available[name] = exact
        except FileNotFoundError:
            font_available[name] = False

    usable: list[tuple[int, Any, int]] = []
    for index, layout in enumerate(prs.slide_layouts):
        count = sum(
            1
            for ph in layout.placeholders
            if ph.placeholder_format.type
            not in {PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER}
            and int(ph.width) > 0
            and int(ph.height) > 0
        )
        if count:
            usable.append((index, layout, count))

    blank_index = _blank_layout_index(prs)
    mapped: dict[SlideTypeId, SlideTypeProfile] = {}
    if len(usable) >= 4:
        for index, layout, count in usable:
            type_id = _guess_type(layout.name, count)
            if not type_id or type_id in mapped:
                continue
            slots = _layout_slots(layout, type_id, minor)
            if not _covers_type(
                slots, type_id, int(prs.slide_width), int(prs.slide_height), minor
            ) or not _geometry_ok(slots, int(prs.slide_width), int(prs.slide_height)):
                continue
            mapped[type_id] = SlideTypeProfile(
                id=type_id,
                label=LABELS[type_id],
                layout_index=index,
                slots=slots,
                supports_chart=type_id == "chart_insight",
                origin="layout",
            )
    else:
        mapped.update(_cluster_slide_profiles(prs, minor, blank_index))

    unmapped = [type_id for type_id in REQUIRED_TYPES if type_id not in mapped]
    for type_id in unmapped:
        mapped[type_id] = SlideTypeProfile(
            id=type_id,
            label=LABELS[type_id],
            layout_index=blank_index,
            slots=_component_slots(type_id, int(prs.slide_width), int(prs.slide_height), minor),
            supports_chart=type_id == "chart_insight",
            origin="component",
        )

    return TemplateProfile(
        source=path.name,
        slide_size=SlideSize(w_emu=int(prs.slide_width), h_emu=int(prs.slide_height)),
        theme=ThemeProfile(
            major_font=major,
            minor_font=minor,
            font_available=font_available,
            colors=colors,
        ),
        slide_types=[mapped[type_id] for type_id in REQUIRED_TYPES],
        missing_types=unmapped,
    )


def compact_profile(profile: TemplateProfile) -> dict[str, Any]:
    return {
        "source": profile.source,
        "slide_size": {"w_emu": profile.slide_size.w_emu, "h_emu": profile.slide_size.h_emu},
        "theme": {
            "major_font": profile.theme.major_font,
            "minor_font": profile.theme.minor_font,
            "colors": profile.theme.colors,
        },
        "slide_types": [
            {
                "id": slide.id,
                "label": slide.label,
                "slots": [
                    {"name": slot.name, "max_chars": slot.max_chars, "font_pt": slot.font_pt}
                    for slot in slide.slots
                ],
            }
            for slide in profile.slide_types
        ],
        "missing_types": profile.missing_types,
    }
