from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from pptx.enum.dml import MSO_COLOR_TYPE, MSO_FILL_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

from .capacity import text_fits
from .profile import open_presentation


def LintIssue(**kwargs: Any) -> dict[str, Any]:
    return dict(kwargs)

EMU = 914400
FORBIDDEN = re.compile(r"lorem|ipsum|\bx{3,}\b|\[insert|TODO", re.IGNORECASE)


def _bbox(shape: Any, pad: int = 0) -> tuple[int, int, int, int]:
    return (
        int(shape.left) - pad,
        int(shape.top) - pad,
        int(shape.left + shape.width) + pad,
        int(shape.top + shape.height) + pad,
    )


def _intersects(a: tuple[int, int, int, int], b: tuple[int, int, int, int]) -> bool:
    return min(a[2], b[2]) > max(a[0], b[0]) and min(a[3], b[3]) > max(a[1], b[1])


def _shape_text(shape: Any) -> str:
    if getattr(shape, "has_text_frame", False):
        return shape.text.strip()
    if getattr(shape, "has_table", False):
        return " ".join(cell.text for row in shape.table.rows for cell in row.cells)
    return ""


def _font_sizes(shape: Any) -> list[float]:
    result: list[float] = []
    if not getattr(shape, "has_text_frame", False):
        return result
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.font.size:
                result.append(run.font.size.pt)
    return result


def _luminance(rgb: tuple[int, int, int]) -> float:
    values = []
    for item in rgb:
        channel = item / 255
        values.append(channel / 12.92 if channel <= 0.04045 else ((channel + 0.055) / 1.055) ** 2.4)
    return 0.2126 * values[0] + 0.7152 * values[1] + 0.0722 * values[2]


def contrast_ratio(foreground: tuple[int, int, int], background: tuple[int, int, int]) -> float:
    first, second = sorted((_luminance(foreground), _luminance(background)), reverse=True)
    return (first + 0.05) / (second + 0.05)


def _tuple_color(value: Any) -> tuple[int, int, int] | None:
    if value is None:
        return None
    clean = str(value)
    if len(clean) != 6:
        return None
    try:
        result = tuple(int(clean[index : index + 2], 16) for index in (0, 2, 4))
        return result[0], result[1], result[2]
    except ValueError:
        return None


def _shape_contrast(shape: Any) -> float | None:
    if not getattr(shape, "has_text_frame", False):
        return None
    background = (255, 255, 255)
    try:
        if (
            shape.fill.type == MSO_FILL_TYPE.SOLID
            and shape.fill.fore_color.type == MSO_COLOR_TYPE.RGB
        ):
            background = _tuple_color(shape.fill.fore_color.rgb) or background
    except (AttributeError, TypeError):
        pass
    ratios: list[float] = []
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            try:
                if run.font.color.type == MSO_COLOR_TYPE.RGB:
                    foreground = _tuple_color(run.font.color.rgb)
                    if foreground:
                        ratios.append(contrast_ratio(foreground, background))
            except (AttributeError, TypeError):
                continue
    return min(ratios) if ratios else None


def _text_overflow(shape: Any, fallback_font: str) -> bool:
    if not getattr(shape, "has_text_frame", False) or not shape.text.strip():
        return False
    sizes = _font_sizes(shape)
    size = min(sizes) if sizes else 18.0
    typeface = fallback_font
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.font.name and not run.font.name.startswith("+"):
                typeface = run.font.name
                break
    return not text_fits(
        shape.text,
        typeface=typeface,
        font_pt=size,
        width_emu=max(1, int(shape.width)),
        height_emu=max(1, int(shape.height)),
    )


def lint_deck(path: Path, fallback_font: str = "Arial") -> list[dict[str, Any]]:
    prs = open_presentation(path)
    issues: list[dict[str, Any]] = []
    margin = int(0.4 * EMU)
    gap_pad = int(0.075 * EMU)
    for slide_index, slide in enumerate(prs.slides, start=1):
        relevant = []
        for shape in slide.shapes:
            name = getattr(shape, "name", "shape")
            if shape.shape_type == MSO_SHAPE_TYPE.LINE:
                continue
            box = _bbox(shape)
            if box[0] < 0 or box[1] < 0 or box[2] > prs.slide_width or box[3] > prs.slide_height:
                issues.append(LintIssue(slide=slide_index, code="OUT_OF_BOUNDS", severity="error", message="Shape extends outside the slide", shape=name))
            is_full_bleed = int(shape.width) >= int(prs.slide_width * 0.95) and int(shape.height) >= int(prs.slide_height * 0.95)
            is_page_number = name == "DA_page_number"
            if not is_full_bleed and not is_page_number and (
                box[0] < margin or box[1] < margin or prs.slide_width - box[2] < margin or prs.slide_height - box[3] < margin
            ):
                issues.append(LintIssue(slide=slide_index, code="EDGE_MARGIN", severity="error", message="Shape breaches the 0.4 inch safe margin", shape=name))
            text = _shape_text(shape)
            if text and FORBIDDEN.search(text):
                issues.append(LintIssue(slide=slide_index, code="PLACEHOLDER_TEXT", severity="error", message="Unresolved placeholder text", shape=name))
            sizes = _font_sizes(shape)
            if sizes and min(sizes) < 10:
                issues.append(LintIssue(slide=slide_index, code="FONT_TOO_SMALL", severity="error", message=f"Effective font is {min(sizes):.1f}pt", shape=name))
            ratio = _shape_contrast(shape)
            if ratio is not None and ratio < 4.5:
                issues.append(
                    LintIssue(
                        slide=slide_index,
                        code="LOW_CONTRAST",
                        severity="error",
                        message=f"Text contrast is {ratio:.2f}:1; minimum is 4.5:1",
                        shape=name,
                    )
                )
            if _text_overflow(shape, fallback_font):
                issues.append(LintIssue(slide=slide_index, code="TEXT_OVERFLOW", severity="error", message="Measured text exceeds its box", shape=name))
            if getattr(shape, "is_placeholder", False) and shape.placeholder_format.type not in {PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER} and not text:
                issues.append(LintIssue(slide=slide_index, code="EMPTY_PLACEHOLDER", severity="error", message="Empty placeholder remains on the slide", shape=name))
            if name.startswith("DA_") and name != "DA_page_number":
                relevant.append(shape)
        for index, first in enumerate(relevant):
            for second in relevant[index + 1 :]:
                if _intersects(_bbox(first, gap_pad), _bbox(second, gap_pad)):
                    issues.append(LintIssue(slide=slide_index, code="ELEMENT_GAP", severity="error", message="Elements overlap or are closer than 0.15 inches", shape=f"{first.name} / {second.name}"))
    return issues
