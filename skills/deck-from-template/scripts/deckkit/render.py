from __future__ import annotations

import json
import logging
import shutil
import subprocess
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Pt

from .capacity import text_fits
from .models import CHART_TYPES, ContentPlan, ContentSlide, TemplateProfile
from .profile import open_presentation

EMU = 914400
logger = logging.getLogger(__name__)


class RenderError(RuntimeError):
    pass


def _rgb(value: str) -> RGBColor:
    clean = value.lstrip("#")
    return RGBColor.from_string(clean if len(clean) == 6 else "171717")


def _luminance(value: str) -> float:
    clean = value.lstrip("#")
    channels = [int(clean[index : index + 2], 16) / 255 for index in (0, 2, 4)]
    linear = [
        item / 12.92 if item <= 0.04045 else ((item + 0.055) / 1.055) ** 2.4
        for item in channels
    ]
    return 0.2126 * linear[0] + 0.7152 * linear[1] + 0.0722 * linear[2]


def _contrast(first: str, second: str) -> float:
    high, low = sorted((_luminance(first), _luminance(second)), reverse=True)
    return (high + 0.05) / (low + 0.05)


def _best_text(fill: str, dark: str, light: str) -> str:
    return dark if _contrast(fill, dark) >= _contrast(fill, light) else light


def _remove_all_slides(prs: Presentation) -> None:
    slide_ids = prs.slides._sldIdLst  # type: ignore[attr-defined]
    for slide_id in list(slide_ids):
        prs.part.drop_rel(slide_id.rId)
        slide_ids.remove(slide_id)


def _set_shape_name(shape: Any, name: str) -> None:
    shape.name = f"DA_{name}"


def _set_run(run: Any, text: str, font_pt: float, color: str, *, major: bool = False) -> None:
    run.text = text
    run.font.name = "+mj-lt" if major else "+mn-lt"
    run.font.size = Pt(font_pt)
    run.font.color.rgb = _rgb(color)


def _text_box(
    slide: Any,
    name: str,
    text: str,
    box: tuple[int, int, int, int],
    font_pt: float,
    color: str,
    *,
    major: bool = False,
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    valign: MSO_ANCHOR = MSO_ANCHOR.TOP,
) -> Any:
    shape = slide.shapes.add_textbox(*box)
    _set_shape_name(shape, name)
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = None
    frame.vertical_anchor = valign
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    _set_run(run, text, font_pt, color, major=major)
    run.font.bold = bold
    return shape


def _line(slide: Any, name: str, x: int, y: int, w: int, color: str, width_pt: float = 2) -> Any:
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x, y, x + w, y)
    _set_shape_name(shape, name)
    shape.line.color.rgb = _rgb(color)
    shape.line.width = Pt(width_pt)
    return shape


def _value(slots: dict[str, Any], key: str, default: str = "") -> str:
    value = slots.get(key, default)
    return value if isinstance(value, str) else str(value)


def _validate_slot_text(slide_type: Any, slots: dict[str, Any], font: str) -> None:
    known = {slot.name for slot in slide_type.slots}
    unknown = [name for name in slots if name not in known]
    if unknown:
        raise RenderError(
            f"{slide_type.id} has no slot named {', '.join(sorted(unknown))}; "
            f"valid slots are {', '.join(sorted(known))}"
        )
    for slot in slide_type.slots:
        value = slots.get(slot.name)
        if not isinstance(value, str) or not value:
            continue
        if len(value) > slot.max_chars:
            raise RenderError(
                f"{slide_type.id}.{slot.name} has {len(value)} characters; cap is {slot.max_chars}"
            )
        if not any(
            text_fits(
                value,
                typeface=font,
                font_pt=size,
                width_emu=slot.width,
                height_emu=slot.height,
            )
            for size in (slot.font_steps or [slot.font_pt])
        ):
            raise RenderError(f"{slide_type.id}.{slot.name} does not fit its measured slot")


def _fit_font(slot: Any, text: str, font: str) -> float:
    for size in slot.font_steps or [slot.font_pt]:
        if text_fits(
            text,
            typeface=font,
            font_pt=size,
            width_emu=slot.width,
            height_emu=slot.height,
        ):
            return float(size)
    raise RenderError(f"{slot.name} does not fit")


class _Chart:
    def __init__(self, value: Any) -> None:
        if not isinstance(value, dict):
            raise RenderError("chart slot must be an object")
        missing = [key for key in ("type", "categories", "series", "title") if key not in value]
        if missing:
            raise RenderError(f"chart is missing {', '.join(missing)}")
        self.type = str(value["type"])
        if self.type not in CHART_TYPES:
            raise RenderError(f"unsupported chart type {self.type!r}")
        self.categories = [str(item) for item in value["categories"]]
        self.title = str(value["title"])
        self.unit = str(value.get("unit", ""))
        self.source = str(value.get("source", ""))
        if len(self.title.split()) < 3:
            raise RenderError("chart title must state a claim, not a label")
        raw_series = value["series"]
        if not raw_series:
            raise RenderError("chart needs at least one series")
        self.series = [_Series(item) for item in raw_series]
        if self.type != "scatter":
            for series in self.series:
                if len(series.values) != len(self.categories):
                    raise RenderError("every chart series must align with categories")


class _Series:
    def __init__(self, value: Any) -> None:
        if not isinstance(value, dict) or "name" not in value or "values" not in value:
            raise RenderError("each chart series needs {name, values}")
        self.name = str(value["name"])
        self.values = [None if item is None else float(item) for item in value["values"]]


def _add_chart(slide: Any, spec_value: dict[str, Any], box: tuple[int, int, int, int], colors: dict[str, str]) -> Any:
    spec = _Chart(spec_value)
    if spec.type == "scatter":
        chart_data = XyChartData()
        xs = []
        for index, category in enumerate(spec.categories):
            try:
                xs.append(float(category))
            except ValueError:
                xs.append(float(index + 1))
        for series in spec.series:
            added = chart_data.add_series(series.name)
            for x, y in zip(xs, series.values):
                if y is not None:
                    added.add_data_point(x, y)
        chart_type = XL_CHART_TYPE.XY_SCATTER
    else:
        chart_data = CategoryChartData()
        chart_data.categories = spec.categories
        if spec.type == "waterfall":
            values = [float(value or 0) for value in spec.series[0].values]
            bases: list[float] = []
            positive: list[float] = []
            negative: list[float] = []
            running = 0.0
            for value in values:
                if value >= 0:
                    bases.append(running)
                    positive.append(value)
                    negative.append(0)
                else:
                    bases.append(running + value)
                    positive.append(0)
                    negative.append(abs(value))
                running += value
            chart_data.add_series("Base", bases)
            chart_data.add_series("Increase", positive)
            chart_data.add_series("Decrease", negative)
            chart_type = XL_CHART_TYPE.COLUMN_STACKED
        else:
            for series in spec.series:
                chart_data.add_series(series.name, series.values)
            chart_type = {
                "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
                "bar": XL_CHART_TYPE.BAR_CLUSTERED,
                "line": XL_CHART_TYPE.LINE_MARKERS,
                "stacked_bar": XL_CHART_TYPE.BAR_STACKED,
            }[spec.type]
    frame = slide.shapes.add_chart(chart_type, *box, chart_data)
    _set_shape_name(frame, "native_chart")
    chart = frame.chart
    chart.has_title = True
    chart.chart_title.text_frame.text = spec.title
    chart.chart_title.text_frame.paragraphs[0].font.size = Pt(12)
    chart.has_legend = len(chart.series) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    palette = [colors.get(f"accent{i}", colors.get("accent1", "2563EB")) for i in range(1, 7)]
    for index, series in enumerate(chart.series):
        if spec.type == "waterfall" and index == 0:
            series.format.fill.background()
            series.format.line.fill.background()
        else:
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = _rgb(palette[index % len(palette)])
    return frame


def _add_table(slide: Any, value: Any, box: tuple[int, int, int, int], colors: dict[str, str]) -> Any:
    if not isinstance(value, dict):
        raise RenderError("table slot must contain {headers, rows}")
    headers = [str(item) for item in value.get("headers", [])]
    rows = [[str(cell) for cell in row] for row in value.get("rows", [])]
    if not headers or not rows or any(len(row) != len(headers) for row in rows):
        raise RenderError("table rows must align with headers")
    if len(rows) > 8 or len(headers) > 7:
        raise RenderError("table capacity is 8 rows by 7 columns")
    shape = slide.shapes.add_table(len(rows) + 1, len(headers), *box)
    _set_shape_name(shape, "native_table")
    table = shape.table
    for column_index, header in enumerate(headers):
        cell = table.cell(0, column_index)
        cell.text = header
        cell.fill.solid()
        header_fill = colors.get("dk1", "171717")
        cell.fill.fore_color.rgb = _rgb(header_fill)
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.size = Pt(12)
            paragraph.font.bold = True
            paragraph.font.color.rgb = _rgb(
                _best_text(
                    header_fill,
                    colors.get("dk1", "171717"),
                    colors.get("lt1", "FFFFFF"),
                )
            )
    for row_index, row in enumerate(rows, start=1):
        for column_index, text in enumerate(row):
            cell = table.cell(row_index, column_index)
            cell.text = text
            if row_index % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = _rgb(colors.get("lt2", "F8FAFC"))
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Pt(11)
                paragraph.font.color.rgb = _rgb(colors.get("dk1", "171717"))
    return shape


def _graphviz_positions(flow: dict[str, Any]) -> dict[str, tuple[float, float]] | None:
    dot = shutil.which("dot")
    if not dot:
        logger.warning("Graphviz is unavailable; using deterministic stage-flow fallback")
        return None
    nodes = flow.get("nodes", [])
    edges = flow.get("edges", [])
    lines = ["digraph G {", "rankdir=LR;", "node [shape=box];"]
    for node in nodes:
        lines.append(f'{json.dumps(str(node["id"]))} [label={json.dumps(str(node.get("label", node["id"])))}];')
    for edge in edges:
        lines.append(f'{json.dumps(str(edge["from"]))} -> {json.dumps(str(edge["to"]))};')
    lines.append("}")
    try:
        result = subprocess.run(
            [dot, "-Tjson"],
            input="\n".join(lines),
            text=True,
            capture_output=True,
            timeout=10,
            check=True,
        )
        data = json.loads(result.stdout)
        positions = {}
        for obj in data.get("objects", []):
            if "pos" in obj and "name" in obj:
                x, y = (float(number) for number in obj["pos"].split(","))
                positions[str(obj["name"])] = (x, y)
        return positions or None
    except (
        OSError,
        KeyError,
        TypeError,
        ValueError,
        json.JSONDecodeError,
        subprocess.SubprocessError,
    ) as exc:
        logger.warning("Graphviz layout failed; using deterministic stage-flow fallback: %s", exc)
        return None


def _add_flow(slide: Any, flow: Any, box: tuple[int, int, int, int], colors: dict[str, str]) -> None:
    if not isinstance(flow, dict):
        raise RenderError("flow slot must contain {nodes, edges}")
    nodes = flow.get("nodes", [])
    edges = flow.get("edges", [])
    if not 2 <= len(nodes) <= 8:
        raise RenderError("flow supports 2–8 nodes")
    left, top, width, height = box
    layout = flow.get("layout", "stage_flow")
    if layout == "pyramid":
        gap = int(0.16 * EMU)
        level_h = int((height - gap * (len(nodes) - 1)) / len(nodes))
        for index, node in enumerate(nodes):
            fraction = (index + 1) / len(nodes)
            level_w = int(width * (0.42 + 0.52 * fraction))
            x = left + (width - level_w) // 2
            y = top + index * (level_h + gap)
            shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.TRAPEZOID, x, y, level_w, level_h)
            _set_shape_name(shape, f"pyramid_level_{index + 1}")
            shape.fill.solid()
            level_fill = colors.get(
                f"accent{(index % 6) + 1}", colors.get("accent1", "2563EB")
            )
            shape.fill.fore_color.rgb = _rgb(level_fill)
            shape.line.fill.background()
            frame = shape.text_frame
            frame.clear()
            frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            paragraph = frame.paragraphs[0]
            paragraph.alignment = PP_ALIGN.CENTER
            run = paragraph.add_run()
            _set_run(
                run,
                str(node.get("label", node["id"])),
                13,
                _best_text(
                    level_fill,
                    colors.get("dk1", "171717"),
                    colors.get("lt1", "FFFFFF"),
                ),
            )
            run.font.bold = True
        return
    positions = None if layout == "two_by_two" else _graphviz_positions(flow)
    if positions:
        xs = [item[0] for item in positions.values()]
        ys = [item[1] for item in positions.values()]
        min_x, max_x = min(xs), max(xs)
        min_y, max_y = min(ys), max(ys)
        normalised = {
            key: (
                (x - min_x) / max(1, max_x - min_x),
                1 - (y - min_y) / max(1, max_y - min_y),
            )
            for key, (x, y) in positions.items()
        }
    elif layout == "two_by_two" and len(nodes) == 4:
        normalised = {
            str(node["id"]): (float(index % 2), float(index // 2))
            for index, node in enumerate(nodes)
        }
    else:
        normalised = {
            str(node["id"]): (index / max(1, len(nodes) - 1), 0.5)
            for index, node in enumerate(nodes)
        }
    node_w = min(int(width * 0.18), int(1.85 * EMU))
    node_h = min(int(height * 0.28), int(0.85 * EMU))
    centers: dict[str, tuple[int, int]] = {}
    for node in nodes:
        nx, ny = normalised.get(str(node["id"]), (0.5, 0.5))
        centers[str(node["id"])] = (
            left + int(node_w / 2 + nx * (width - node_w)),
            top + int(node_h / 2 + ny * (height - node_h)),
        )
    for index, edge in enumerate(edges):
        start = centers.get(str(edge.get("from")))
        end = centers.get(str(edge.get("to")))
        if not start or not end:
            continue
        connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, *start, *end)
        _set_shape_name(connector, f"flow_edge_{index + 1}")
        connector.line.color.rgb = _rgb(colors.get("accent1", "2563EB"))
        connector.line.width = Pt(1.5)
    for index, node in enumerate(nodes):
        center = centers[str(node["id"])]
        x, y = center[0] - node_w // 2, center[1] - node_h // 2
        shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, node_w, node_h)
        _set_shape_name(shape, f"flow_node_{index + 1}")
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(colors.get("lt2", "F8FAFC"))
        shape.line.color.rgb = _rgb(colors.get("accent1", "2563EB"))
        frame = shape.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        paragraph = frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.CENTER
        run = paragraph.add_run()
        _set_run(run, str(node.get("label", node["id"])), 13, colors.get("dk1", "171717"))
        run.font.bold = True


def _render_component(slide: Any, content: ContentSlide, type_profile: Any, profile: TemplateProfile) -> None:
    slots = content.slots
    slot_map = {slot.name: slot for slot in type_profile.slots}
    colors = profile.theme.colors
    dark = colors.get("dk1", "171717")
    muted = colors.get("dk2", "334155")
    accent = colors.get("accent1", "2563EB")
    font = profile.theme.minor_font

    for slot in type_profile.slots:
        if (
            slot.name in {"chart", "table", "flow"}
            or (content.slide_type == "big_numbers" and slot.name.startswith(("number_", "label_")))
            or slot.name not in slots
        ):
            continue
        text = _value(slots, slot.name)
        size = _fit_font(slot, text, font)
        is_heading = slot.name in {"headline", "number_1", "number_2", "number_3", "number_4"}
        _text_box(
            slide,
            slot.name,
            text,
            (slot.left, slot.top, slot.width, slot.height),
            size,
            dark if slot.name != "source" else muted,
            major=is_heading,
            bold=is_heading or slot.name.endswith("_title"),
        )

    if content.slide_type == "big_numbers":
        present = [
            index
            for index in range(1, 5)
            if _value(slots, f"number_{index}") or _value(slots, f"label_{index}")
        ]
        count = max(1, len(present))
        margin = int(profile.slide_size.w_emu * 0.065)
        content_w = profile.slide_size.w_emu - 2 * margin
        gap = int(profile.slide_size.w_emu * 0.018)
        col_w = int((content_w - gap * (count - 1)) / count)
        for position, index in enumerate(present):
            x = margin + position * (col_w + gap)
            number_slot = slot_map[f"number_{index}"]
            label_slot = slot_map[f"label_{index}"]
            _text_box(
                slide,
                f"number_{index}",
                _value(slots, f"number_{index}"),
                (x, number_slot.top, col_w, number_slot.height),
                number_slot.font_pt,
                dark,
                major=True,
                bold=True,
            )
            _text_box(
                slide,
                f"label_{index}",
                _value(slots, f"label_{index}"),
                (x, label_slot.top, col_w, label_slot.height),
                label_slot.font_pt,
                muted,
            )

    if "headline" in slot_map:
        headline = slot_map["headline"]
        _line(slide, "headline_rule", headline.left, headline.top - int(0.12 * EMU), int(0.52 * EMU), accent, 3)
    if "chart" in slot_map and "chart" in slots:
        slot = slot_map["chart"]
        _add_chart(slide, slots["chart"], (slot.left, slot.top, slot.width, slot.height), colors)
    if "table" in slot_map and "table" in slots:
        slot = slot_map["table"]
        _add_table(slide, slots["table"], (slot.left, slot.top, slot.width, slot.height), colors)
    if "flow" in slot_map and "flow" in slots:
        slot = slot_map["flow"]
        _add_flow(slide, slots["flow"], (slot.left, slot.top, slot.width, slot.height), colors)

    # Minimal page marker, outside the content field but inside the safe margin.
    _text_box(
        slide,
        "page_number",
        f"{content.n:02d}",
        (int(profile.slide_size.w_emu * 0.91), int(profile.slide_size.h_emu * 0.91), int(0.35 * EMU), int(0.18 * EMU)),
        10,
        muted,
        align=PP_ALIGN.RIGHT,
    )


def _find_placeholder(slide: Any, ph_idx: int) -> Any | None:
    for shape in slide.placeholders:
        if int(shape.placeholder_format.idx) == ph_idx:
            return shape
    return None


def _fill_placeholder(shape: Any, text: str, slot: Any, profile: TemplateProfile) -> None:
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = None
    paragraph = frame.paragraphs[0]
    run = paragraph.add_run()
    size = _fit_font(slot, text, profile.theme.minor_font)
    _set_run(
        run,
        text,
        size,
        slot.font_color or profile.theme.colors.get("dk1", "171717"),
        major=slot.name == "headline",
    )
    if slot.font_weight == "bold" or slot.name == "headline" or slot.name.endswith("_title"):
        run.font.bold = True
    _set_shape_name(shape, slot.name)


def _remove_unused_placeholders(slide: Any) -> None:
    for shape in list(slide.placeholders):
        kind = shape.placeholder_format.type
        if kind in {PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER}:
            continue
        if not shape.has_text_frame or not shape.text.strip():
            shape._element.getparent().remove(shape._element)


def _render_layout(slide: Any, content: ContentSlide, type_profile: Any, profile: TemplateProfile) -> None:
    for slot in type_profile.slots:
        value = content.slots.get(slot.name)
        if value is None:
            continue
        if slot.name == "chart":
            _add_chart(slide, value, (slot.left, slot.top, slot.width, slot.height), profile.theme.colors)
        elif slot.name == "table":
            _add_table(slide, value, (slot.left, slot.top, slot.width, slot.height), profile.theme.colors)
        elif slot.name == "flow":
            _add_flow(slide, value, (slot.left, slot.top, slot.width, slot.height), profile.theme.colors)
        else:
            placeholder = (
                _find_placeholder(slide, slot.ph_idx) if slot.ph_idx is not None else None
            )
            if placeholder is None:
                raise RenderError(
                    f"slide {content.n}: slot {slot.name!r} has content but no placeholder "
                    "in the template layout"
                )
            _fill_placeholder(placeholder, str(value), slot, profile)
    _remove_unused_placeholders(slide)


def render_one(prs: Any, content: ContentSlide, profile: TemplateProfile) -> list[dict[str, Any]]:
    """Compose one generative slide onto the end of `prs`. Returns any fallbacks used."""
    fallbacks: list[dict[str, Any]] = []
    type_profile = profile.slide_type(content.slide_type)
    if content.slide_type == "flow" and isinstance(content.slots.get("flow"), dict):
        layout = content.slots["flow"].get("layout", "stage_flow")
        if layout == "stage_flow" and not shutil.which("dot"):
            fallbacks.append(
                {
                    "slide": content.n,
                    "component": "flow",
                    "reason": "Graphviz unavailable; deterministic native stage layout used",
                }
            )
    _validate_slot_text(type_profile, content.slots, profile.theme.minor_font)
    if type_profile.layout_index >= len(prs.slide_layouts):
        raise RenderError(f"layout index {type_profile.layout_index} is missing")
    slide = prs.slides.add_slide(prs.slide_layouts[type_profile.layout_index])
    if type_profile.origin == "layout":
        _render_layout(slide, content, type_profile, profile)
    else:
        _remove_unused_placeholders(slide)
        _render_component(slide, content, type_profile, profile)
    return fallbacks


def render_deck(
    template_path: Path, profile: TemplateProfile, content_plan: ContentPlan, output_path: Path
) -> Path:
    prs = open_presentation(template_path)
    if int(prs.slide_width) != profile.slide_size.w_emu or int(prs.slide_height) != profile.slide_size.h_emu:
        raise RenderError("template profile slide size does not match the template")
    _remove_all_slides(prs)
    fallbacks: list[dict[str, Any]] = []
    for expected_n, content in enumerate(content_plan.slides, start=1):
        if content.n != expected_n:
            raise RenderError("content slides must be sequentially numbered")
        fallbacks.extend(render_one(prs, content, profile))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)
    (output_path.parent / "render_fallbacks.json").write_text(
        json.dumps({"count": len(fallbacks), "fallbacks": fallbacks}, indent=2) + "\n",
        encoding="utf-8",
    )
    return output_path
