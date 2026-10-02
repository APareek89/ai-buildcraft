"""Read the reference deck's own slides as reusable archetypes.

Layout extraction (profile.py) answers "what layouts exist". This answers the
harder question: "what is each real slide in this deck *for*, and how do I put
new content into that exact composition". Python does the measuring and the
copying; the semantic judgement — this slide is a feature showcase, that one is
a competitive 2x2 — is made by the model from the rendered thumbnails.
"""

from __future__ import annotations

import copy
import re
from pathlib import Path
from typing import Any

from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.util import Emu, Pt

from .capacity import measured_capacity, text_fits
from .profile import _theme_from_package, open_presentation

EMU = 914400
DECORATIVE = {PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER}


# ---------------------------------------------------------------- fingerprints


def _kind(shape: Any) -> str:
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return "picture"
    if getattr(shape, "has_table", False):
        return "table"
    if getattr(shape, "has_chart", False):
        return "chart"
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return "group"
    if shape.shape_type == MSO_SHAPE_TYPE.LINE:
        return "line"
    if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
        return "text"
    if getattr(shape, "has_text_frame", False):
        return "empty_text"
    return "shape"


def _run_style(shape: Any) -> dict[str, Any]:
    if not getattr(shape, "has_text_frame", False):
        return {}
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            colour = None
            try:
                colour = str(run.font.color.rgb) if run.font.color.rgb else None
            except (AttributeError, TypeError):
                colour = None
            return {
                "font_pt": round(run.font.size.pt, 1) if run.font.size else None,
                "bold": bool(run.font.bold),
                "colour": colour,
            }
    return {}


def _text_shape_of(shape: Any) -> str:
    if getattr(shape, "has_table", False):
        table = shape.table
        return f"table {len(table.rows)}x{len(table.columns)}"
    if getattr(shape, "has_text_frame", False):
        return re.sub(r"\s+", " ", shape.text_frame.text).strip()
    return ""


def fingerprint_slide(prs: Any, slide: Any, index: int) -> dict[str, Any]:
    width, height = int(prs.slide_width), int(prs.slide_height)
    shapes: list[dict[str, Any]] = []
    for order, shape in enumerate(slide.shapes):
        if (
            getattr(shape, "is_placeholder", False)
            and shape.placeholder_format.type in DECORATIVE
        ):
            continue
        text = _text_shape_of(shape)
        entry: dict[str, Any] = {
            "shape": order,
            "name": shape.name,
            "kind": _kind(shape),
            "box": [
                round(int(shape.left) / width, 4),
                round(int(shape.top) / height, 4),
                round(int(shape.width) / width, 4),
                round(int(shape.height) / height, 4),
            ],
            "emu": [int(shape.left), int(shape.top), int(shape.width), int(shape.height)],
            "text": text[:160],
            "chars": len(text),
        }
        style = _run_style(shape)
        if style:
            entry["style"] = style
        if getattr(shape, "has_table", False):
            entry["table"] = {
                "rows": len(shape.table.rows),
                "cols": len(shape.table.columns),
                "header": [cell.text.strip()[:40] for cell in shape.table.rows[0].cells],
            }
        shapes.append(entry)
    return {
        "slide": index,
        "layout": slide.slide_layout.name,
        "shape_count": len(shapes),
        "repeating_groups": detect_repetition(shapes),
        "shapes": shapes,
    }


def detect_repetition(shapes: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Find n-up rhythms — the tell that a slide is a card row or a stage strip.

    Shapes of near-identical size, evenly spaced on one axis, are the parallel
    slots an archetype exposes (three feature cards, five process arrows).
    """
    groups: list[dict[str, Any]] = []
    buckets: dict[tuple[float, float], list[dict[str, Any]]] = {}
    for entry in shapes:
        if entry["kind"] in {"line", "group"}:
            continue
        key = (round(entry["box"][2], 2), round(entry["box"][3], 2))
        buckets.setdefault(key, []).append(entry)
    for (bw, bh), members in buckets.items():
        if len(members) < 2:
            continue
        for axis, index in (("row", 0), ("column", 1)):
            ordered = sorted(members, key=lambda item: item["box"][index])
            positions = [item["box"][index] for item in ordered]
            other = {round(item["box"][1 - index], 2) for item in ordered}
            if len(other) != 1:
                continue
            gaps = [
                round(positions[i + 1] - positions[i], 3) for i in range(len(positions) - 1)
            ]
            if gaps and max(gaps) - min(gaps) <= 0.01:
                groups.append(
                    {
                        "axis": axis,
                        "count": len(ordered),
                        "size": [bw, bh],
                        "shapes": [item["shape"] for item in ordered],
                        "texts": [item["text"][:40] for item in ordered],
                    }
                )
            break
    return groups


def fingerprint_deck(path: Path) -> dict[str, Any]:
    prs = open_presentation(path)
    major, minor, colours = _theme_from_package(path)
    slides = [
        fingerprint_slide(prs, slide, index)
        for index, slide in enumerate(prs.slides, start=1)
    ]
    return {
        "source": path.name,
        "slide_size": {"w_emu": int(prs.slide_width), "h_emu": int(prs.slide_height)},
        "theme": {"major_font": major, "minor_font": minor, "colors": colours},
        "slide_count": len(slides),
        "slides": slides,
    }


# ------------------------------------------------------------------- measuring


def _shape_at(slide: Any, order: int) -> Any:
    visible = [
        shape
        for shape in slide.shapes
        if not (
            getattr(shape, "is_placeholder", False)
            and shape.placeholder_format.type in DECORATIVE
        )
    ]
    if not 0 <= order < len(visible):
        raise KeyError(f"shape index {order} is out of range for this slide")
    return visible[order]


def measure_catalog(path: Path, catalog: dict[str, Any]) -> dict[str, Any]:
    """Attach measured capacities and resolved geometry to a hand-authored catalog."""
    prs = open_presentation(path)
    _, minor, _ = _theme_from_package(path)
    slides = list(prs.slides)
    measured = copy.deepcopy(catalog)
    for archetype in measured.get("archetypes", []):
        number = int(archetype["source_slide"])
        if not 1 <= number <= len(slides):
            raise ValueError(
                f"{archetype['id']}: source_slide {number} does not exist "
                f"(deck has {len(slides)} slides)"
            )
        slide = slides[number - 1]
        seen: set[str] = set()
        for slot in archetype.get("slots", []):
            name = slot["name"]
            if name in seen:
                raise ValueError(f"{archetype['id']}: duplicate slot {name!r}")
            seen.add(name)
            shape = _shape_at(slide, int(slot["shape"]))
            slot["kind"] = _kind(shape)
            slot["emu"] = [
                int(shape.left),
                int(shape.top),
                int(shape.width),
                int(shape.height),
            ]
            slot["sample"] = _text_shape_of(shape)[:120]
            if slot["kind"] == "table":
                slot["rows"] = len(shape.table.rows)
                slot["cols"] = len(shape.table.columns)
                if slot["rows"] != 1 or slot["cols"] != 1:
                    continue
                # A 1x1 table is a tinted content box; it takes text, so measure it.
            if slot["kind"] == "picture":
                continue
            if slot["kind"] == "table":
                cell = shape.table.cell(0, 0)
                font_pt = next(
                    (
                        run.font.size.pt
                        for paragraph in cell.text_frame.paragraphs
                        for run in paragraph.runs
                        if run.font.size
                    ),
                    14.0,
                )
                slot["sample"] = cell.text_frame.text.strip()[:120]
            else:
                font_pt = _run_style(shape).get("font_pt") or 18.0
            slot["font_pt"] = font_pt
            slot["max_chars"] = measured_capacity(
                typeface=minor,
                font_pt=font_pt,
                width_emu=max(1, int(shape.width)),
                height_emu=max(1, int(shape.height)),
            )
    measured["minor_font"] = minor
    measured["source"] = path.name
    return measured


# --------------------------------------------------------------- clone + refill


def _remap_rels(element: Any, mapping: dict[str, str]) -> None:
    """Point copied XML at the relationship ids it was given in the new part."""
    namespace = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    for node in element.iter():
        for attribute in ("embed", "link", "id", "pict", "dm", "lo", "qs", "cs"):
            key = namespace + attribute
            value = node.get(key)
            if value and value in mapping:
                node.set(key, mapping[value])


def clone_slide(source_slide: Any, dest_prs: Any) -> Any:
    layout = source_slide.slide_layout
    for candidate in dest_prs.slide_layouts:
        if candidate.name == layout.name:
            layout = candidate
            break
    dest = dest_prs.slides.add_slide(layout)
    for shape in list(dest.shapes):
        shape._element.getparent().remove(shape._element)

    mapping: dict[str, str] = {}
    for rid, rel in source_slide.part.rels.items():
        if rel.reltype.endswith("slideLayout"):
            continue
        if rel.is_external:
            mapping[rid] = dest.part.rels._add_relationship(
                rel.reltype, rel.target_ref, is_external=True
            )
        else:
            mapping[rid] = dest.part.rels._add_relationship(rel.reltype, rel.target_part)

    tree = dest.shapes._spTree
    for shape in source_slide.shapes:
        element = copy.deepcopy(shape._element)
        _remap_rels(element, mapping)
        tree.append(element)
    return dest


def _fit_size(text: str, font: str, base: float, width: int, height: int) -> float:
    for size in (base, base - 1, base - 2, base - 3, base - 4):
        if size < 9:
            break
        if text_fits(
            text, typeface=font, font_pt=size, width_emu=width, height_emu=height
        ):
            return size
    return max(9.0, base - 4)


def _dominant_bold(runs: list[Any]) -> bool:
    """The weight that carries most of a mixed paragraph's characters.

    A source line like "**Understand** business context and the expected outcome"
    is mostly body text, so new content written into it should be body weight.
    """
    weight: dict[bool, int] = {}
    for run in runs:
        weight[bool(run.font.bold)] = weight.get(bool(run.font.bold), 0) + len(run.text)
    return max(weight.items(), key=lambda item: item[1])[0] if weight else False


def _fill_text_frame(
    frame: Any, lines: list[str], size: float | None, bold: bool | None = None
) -> None:
    """Write `lines` into `frame`, one per paragraph, preserving existing formatting.

    When the frame already has as many paragraphs as there are lines, each keeps
    its own styling — the 24pt white title and the 14pt subtitle inside one
    pentagon arrow stay distinct. Extra paragraphs are cloned from the last one.
    """
    paragraphs = frame.paragraphs
    if not paragraphs:
        frame.text = lines[0]
        paragraphs = frame.paragraphs
    body = paragraphs[0]._p.getparent()
    template = copy.deepcopy(paragraphs[len(paragraphs) - 1]._p)

    for spare in list(paragraphs[len(lines) :]):
        body.remove(spare._p)
    while len(frame.paragraphs) < len(lines):
        body.append(copy.deepcopy(template))

    for paragraph, text in zip(frame.paragraphs, lines):
        runs = paragraph.runs
        if not runs:
            paragraph.add_run()
            runs = paragraph.runs
        # A source paragraph like "**Understand** business context, ..." has mixed
        # inline emphasis that one run cannot reproduce. Inherit the body run's
        # weight rather than smearing the lead-in's bold across the whole line.
        if bold is not None:
            runs[0].font.bold = bold
        elif len(runs) > 1 and len({bool(run.font.bold) for run in runs}) > 1:
            runs[0].font.bold = _dominant_bold(runs)
        runs[0].text = text
        if size is not None:
            runs[0].font.size = Pt(size)
        for extra in runs[1:]:
            extra._r.getparent().remove(extra._r)


def _as_lines(value: Any) -> list[str]:
    if isinstance(value, list):
        return [str(item) for item in value]
    return str(value).split("\n")


def set_text(
    shape: Any, value: Any, font: str, base_pt: float, bold: bool | None = None
) -> None:
    """Replace a shape's text, keeping the run and paragraph formatting it already has."""
    lines = _as_lines(value)
    size = _fit_size(
        "\n".join(lines), font, base_pt, max(1, int(shape.width)), max(1, int(shape.height))
    )
    _fill_text_frame(shape.text_frame, lines, size if size < base_pt else None, bold)


def set_table(
    shape: Any,
    value: Any,
    font: str = "Arial",
    base_pt: float = 14.0,
    bold: bool | None = None,
) -> None:
    """Overwrite an existing table's text, keeping its exact styling.

    A 1x1 table is a tinted content box, not a grid — decks use them constantly —
    so it accepts a string or a list of bullet lines directly.
    """
    table = shape.table
    columns, rows_available = len(table.columns), len(table.rows)

    if columns == 1 and rows_available == 1 and not isinstance(value, dict):
        lines = _as_lines(value)
        size = _fit_size(
            "\n".join(lines),
            font,
            base_pt,
            max(1, int(shape.width)),
            max(1, int(shape.height)),
        )
        _fill_text_frame(
            table.cell(0, 0).text_frame, lines, size if size < base_pt else None, bold
        )
        return

    if not isinstance(value, dict) or "rows" not in value:
        raise ValueError("table slot needs {rows: [[...]], headers?: [...]}")
    body_rows = [[str(cell) for cell in row] for row in value["rows"]]
    headers = [str(cell) for cell in value.get("headers", [])]
    capacity = rows_available - (1 if headers else 0)
    if len(body_rows) > capacity:
        raise ValueError(
            f"this template table holds {capacity} body rows; {len(body_rows)} were given"
        )
    if any(len(row) != columns for row in body_rows) or (
        headers and len(headers) != columns
    ):
        raise ValueError(f"this template table has {columns} columns")

    offset = 0
    if headers:
        for index, text in enumerate(headers):
            _fill_text_frame(table.cell(0, index).text_frame, [text], None)
        offset = 1
    for row_index in range(capacity):
        for column in range(columns):
            text = body_rows[row_index][column] if row_index < len(body_rows) else ""
            _fill_text_frame(table.cell(row_index + offset, column).text_frame, [text], None)


def set_picture(slide: Any, shape: Any, image_path: Path) -> None:
    """Swap a picture, preserving its exact frame and z-order."""
    if not image_path.exists():
        raise ValueError(f"image not found: {image_path}")
    left, top, width, height = (
        int(shape.left),
        int(shape.top),
        int(shape.width),
        int(shape.height),
    )
    replacement = slide.shapes.add_picture(
        str(image_path), Emu(left), Emu(top), Emu(width), Emu(height)
    )
    old = shape._element
    old.addnext(replacement._element)
    old.getparent().remove(old)


def drop_shape(shape: Any) -> None:
    shape._element.getparent().remove(shape._element)
